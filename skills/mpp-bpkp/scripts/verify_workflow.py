"""Uji dengan data fiktif di folder terpisah. Tidak menyentuh input operasional."""
import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
from zipfile import ZipFile
from lxml import etree as E
import openpyxl
from schema import ROOT, DATA, NARASI, PIA_REQUIRED, VERSION
from model import read_model, compile_plan, InputError, words
from generate_mpp import generate, TOKEN, NS
from office_io import recalculate_xlsx

def fixture(folder,template,label,count,jenis):
    path=folder/(label+'.xlsx');shutil.copy2(template,path);w=openpyxl.load_workbook(path);s=w['DATA']
    key_to_row={}
    for r in range(1,s.max_row+1):
        k=s.cell(r,5).value
        if k and isinstance(k,str) and k.strip():key_to_row[k.strip()]=r
    for r,(key,text,default,required) in DATA.items():
        row_idx=key_to_row.get(key,r)
        c=s.cell(row_idx,2)
        if required and c.value is None:c.value='Data uji sintetis '+key
    for i in range(min(count,3)):
        r=i+4;s.cell(r,2,f'Pegawai Uji {i+1:02}');s.cell(r,3,f'19900101202001{i+1:04}')
    kaper_row=27
    for r in range(4,35):
        if 'Kepala Perwakilan' in str(s.cell(r,1).value or ''):
            kaper_row=r;break
    for i in range(3,count):
        r=i+4
        if r<kaper_row:
            s.cell(r,2,f'Pegawai Uji {i+1:02}');s.cell(r,3,f'19900101202001{i+1:04}')
        elif 'TIM_TAMBAHAN' in w.sheetnames:
            w['TIM_TAMBAHAN'].cell(i-kaper_row+8,1,f'Pegawai Uji {i+1:02}')
            w['TIM_TAMBAHAN'].cell(i-kaper_row+8,2,f'19900101202001{i+1:04}')
    s.cell(kaper_row,2,'Kepala Perwakilan Uji')
    def set_v(k,v):
        if k in key_to_row:s.cell(key_to_row[k],2,v)
    set_v('NAMA_PENUGASAN',label);set_v('TINGKAT_RISIKO','Sedang')
    set_v('TANGGAL_ST',datetime(2026,11,2));set_v('HARI_PERSIAPAN',1);set_v('HARI_PELAKSANAAN',3);set_v('HARI_PENYELESAIAN',1)
    set_v('JENIS_PENUGASAN',jenis);set_v('JALUR_ST','MANUAL');set_v('BENTUK_PENGAWASAN','PIA' if jenis=='PIA' else 'EVALUASI');set_v('SUMBER_PIA','ISIAN')
    if 'LIBUR' in w.sheetnames:
        w['LIBUR']['A5']=datetime(2026,11,5);w['LIBUR']['B5']='Libur sintetis untuk uji kalender'
    for i in range(1,14):
        r=i+4;w['IHP'].cell(r,1,f'IHP-{i:02} Analisis sintetis ketepatan proses nomor {i}')
    w['TEMBUSAN']['A5']='Penerima tembusan sintetis'
    w.save(path);recalculate_xlsx(path)
    m=read_model(path,recalc=False,inspect_only=True)
    names=[p['name'] for p in m['team']]
    plan={'schema_version':VERSION,'input_sha256':m['input_sha256'],'sources_read':[{'path':str(path),'location':'DATA/IHP','used_for':'Pengujian perangkat lunak; bukan dokumen penugasan nyata.'}],'ihp':m['ihp'],'narratives':{k:('Telaah dokumen uji\nKonfirmasi data uji' if k=='METODOLOGI' else 'Narasi sintetis untuk menguji '+k+'.') for k,(v,required) in NARASI.items() if required},'processes':[[f'Tahap sintetis {i}','Tujuan uji','Aktivitas uji','Input uji','Output uji'] for i in range(1,8)],'client_risks':[],'assignment_risks':[{'id':'RP-01','event':'Keterlambatan data uji','cause':'Jadwal penyerahan data uji berubah','impact':'Pengujian sintetis tertunda','controls':'Menetapkan batas penyerahan dan sumber alternatif dalam rencana uji.','owner':names[2],'when':'Persiapan dan pemantauan harian'}],'objectives':[],'tasks':[],'general_split':{'conclusion_share':.1,'final_kk_share':.5}}
    for i,ihp in enumerate(m['ihp'],1):
        rid=f'RK-{i:02}';oid=f'TEO-{i:02}';plan['client_risks'].append({'id':rid,'ihp_ids':[ihp['id']],'event':f'Sasaran sintetis {i} tidak tercapai','cause':'Penyebab potensial sintetis','impact':'Dampak sintetis pada sasaran uji','basis':'Asumsi pengujian perangkat lunak, bukan fakta klien'})
        plan['objectives'].append({'id':oid,'type':'TUJUAN_PIA' if jenis=='PIA' else 'TEO','text':f'Tujuan tentatif sintetis {i}','risk_ids':[rid],'source':{'kind':'USULAN_TENTATIF','basis':'Data uji, tidak berasal dari BPKP pusat'}})
        people=list(dict.fromkeys([names[(i-1)%count],names[(i+12)%count]]))
        plan['tasks'].append({'id':f'UJI-{i:02}','objective_ids':[oid],'objective':f'Menguji rancangan sintetis nomor {i}','steps':[{'text':f'Peroleh data sintetis {i}, bandingkan dengan kriteria uji, dan dokumentasikan hasil uji.','assignees':people,'weight':i%3+1,'reference':f'KKA-UJI-{i:02}'}]})
    if jenis!='PIA':
        pia={k:'Hasil sintetis '+k for k in PIA_REQUIRED}
        for k in ['TANGGAL_ST_PIA','TANGGAL_MULAI_PIA','TANGGAL_SELESAI_PIA','TANGGAL_ND_PIA']:pia[k]='2026-10-01'
        pia['NIP_KT_PIA']='199001012020010003';pia['NIP_KORWAS_PIA']='199001012020010001';pia['HARI_PIA_TERBILANG']='satu';plan['pia']=pia
    plan_file=folder/(label+'.json');plan_file.write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    return m,plan,plan_file

def verify_output(manifest,model):
    folder=Path(manifest['output_folder']);docs=sorted(folder.glob('*.docx'))
    assert len(docs)==(11 if model['jenis']=='PIA' else 12)
    if model['jenis']=='PIA':assert not list(folder.glob('09*'))
    for f in docs:
        with ZipFile(f) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith('.xml'):
                    root=E.fromstring(z.read(n));txt=''.join(root.xpath('//w:t/text()',namespaces=NS));assert not TOKEN.search(txt),(f.name,n)
    for num in [1,3,4,6,8,12]:
        f=next(folder.glob(f'{num:02}_*.docx'))
        with ZipFile(f) as z:
            root=E.fromstring(z.read('word/document.xml'));txt=''.join(root.xpath('//w:t/text()',namespaces=NS))
            for person in model['team']:assert person['name'] in txt,(f.name,person['name'])
    for num in [4,10]:
        with ZipFile(next(folder.glob(f'{num:02}_*.docx'))) as z:
            doc=E.fromstring(z.read('word/document.xml'));txt=''.join(doc.xpath('//w:t/text()',namespaces=NS));assert model['ihp'][-1]['information'] in txt
            with ZipFile(next((ROOT/'assets'/'templates').glob(f'{num:02}_*.docx'))) as original:
                template=E.fromstring(original.read('word/document.xml'))
                assert len(doc.xpath('//w:sectPr',namespaces=NS))==len(template.xpath('//w:sectPr',namespaces=NS)), 'Pengulangan isi mengubah pemisahan section'
    person_hours={p['name']:0 for p in model['team']}
    for step in manifest['traceability']:
        assert step['objective_ids'] and step['risk_ids'] and step['ihp_ids']
        for name,hours in step['hours_by_person'].items():person_hours[name]+=hours
    for person in model['team']:assert abs(person_hours[person['name']]+person['hours'][1]*.1-person['hours'][1])<.001
    return {'docx':len(docs),'pages':sum(r['pages'] for r in manifest['render']),'folder':str(folder)}

def main(root,template):
    root=Path(root);root.mkdir(parents=True,exist_ok=True);results=[]
    original=Path('D:/autoST/KM4_KM5_KM6.xlsx');before=hashlib.sha256(original.read_bytes()).hexdigest()
    for label,count,jenis in [('UJI SINTETIS PIA Nabire',5,'PIA'),('UJI SINTETIS Evaluasi Nabire',17,'PENGAWASAN')]:
        base,plan,plan_file=fixture(root,template,label,count,jenis)
        model=compile_plan(base,plan)
        report=generate(model,root/'hasil',pdf=True);results.append({'scenario':label,'status':'passed',**verify_output(report,model)})
        for description,mutator in [
            ('Risiko terputus dari Insilwas',lambda x:x['client_risks'][0].update(ihp_ids=['IHP-TIDAK-ADA'])),
            ('KM 6 melompati TAO/TEO',lambda x:x['tasks'][0].update(objective_ids=[])),
            ('Pelaksana di luar tim',lambda x:x['tasks'][0]['steps'][0].update(assignees=['Bukan Anggota Tim'])),
            ('Insilwas tanpa risiko',lambda x:x['client_risks'].pop()),
        ]:
            bad=deepcopy(plan);mutator(bad)
            try:compile_plan(deepcopy(base),bad)
            except InputError:results.append({'scenario':description,'status':'correctly_rejected'})
            else:raise AssertionError(description)
    assert words(11)=='sebelas' and words(21)=='dua puluh satu'
    assert hashlib.sha256(original.read_bytes()).hexdigest()==before
    results.append({'scenario':'Workbook asli tidak berubah','status':'passed','sha256':before})
    (root/'hasil_pengujian.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False),flush=True)

if __name__=='__main__':
    import sys
    sys.stdout.reconfigure(encoding='utf-8')
    p=argparse.ArgumentParser();p.add_argument('root',type=Path);p.add_argument('template',type=Path);a=p.parse_args();main(a.root,a.template)
