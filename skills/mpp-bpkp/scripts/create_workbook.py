"""Buat workbook isian mandiri; tidak mengubah workbook autoST asal."""
import argparse
import json
from pathlib import Path
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from schema import ROOT, VERSION, DATA, NARASI, TABLES, BASE_ROLES, PIA_REQUIRED

BLUE='17365D'; YELLOW='FFF2CC'; PALE='E7E6E6'; WHITE='FFFFFF'; MUTED='595959'

def base(ws, title, note, headers):
    ws.sheet_view.showGridLines=False
    ws.append([title]);ws.append([note]);ws.append([]);ws.append(headers)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(2,len(headers)))
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max(2,len(headers)))
    ws.row_dimensions[1].height=28;ws.row_dimensions[2].height=36
    for row in ws:
        for c in row:
            c.font=Font(name='Arial',size=11)
            c.alignment=Alignment(vertical='top',wrap_text=True)
    for r in [1,4]:
        for c in ws[r]:c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name='Arial',size=11,bold=True,color=WHITE)
    ws.freeze_panes='B5';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:4'

def editable(c, value=None, explanation=''):
    c.value=value;c.fill=PatternFill('solid',fgColor=YELLOW)
    edge=Side(style='thin',color='CFD9E8')
    c.border=Border(left=edge,right=edge,top=edge,bottom=edge)
    c.font=Font(name='Arial',size=11,color='0000FF')
    c.alignment=Alignment(vertical='top',wrap_text=True)
    if explanation:c.comment=Comment(explanation,'MPP BPKP')

def calculated(c, formula, explanation=''):
    editable(c,formula,explanation);c.fill=PatternFill('solid',fgColor=PALE)
    c.font=Font(name='Arial',size=11,color='000000')

def guide_cell(c, text):
    edge=Side(style='thin',color='E0E0E0')
    c.value=text;c.border=Border(left=edge,right=edge,top=edge,bottom=edge)
    c.font=Font(name='Arial',size=10,italic=True,color=MUTED)
    c.alignment=Alignment(vertical='top',wrap_text=True)

def dropdown(ws, cells, choices):
    dv=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=True)
    dv.errorTitle='Pilihan tidak sesuai';dv.error='Pilih salah satu nilai dalam daftar.';dv.showErrorMessage=True
    ws.add_data_validation(dv)
    for addr in cells:dv.add(addr)

def create(path, overwrite=False):
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)

    # 1. Load and copy Pegawai sheet (HIDDEN)
    pegawai_src = None
    for candidate in [ROOT/'assets'/'Isian_DATA_MPP_BPKP.xlsx', Path('D:/autoST/KM4_KM5_KM6.xlsx')]:
        if candidate.is_file():
            try:
                src_wb = openpyxl.load_workbook(candidate, data_only=True)
                if 'Pegawai' in src_wb.sheetnames:
                    pegawai_src = src_wb['Pegawai'];break
            except Exception: pass
    ws_peg = wb.create_sheet('Pegawai')
    if pegawai_src:
        for r in range(1, pegawai_src.max_row + 1):
            for c in range(1, pegawai_src.max_column + 1):
                ws_peg.cell(r, c, pegawai_src.cell(r, c).value)
    ws_peg.column_dimensions['A'].width = 38
    ws_peg.column_dimensions['B'].width = 25
    ws_peg.column_dimensions['C'].width = 15
    ws_peg.column_dimensions['D'].width = 13
    ws_peg.column_dimensions['E'].width = 18
    ws_peg.sheet_state = 'hidden'  # HIDE sheet Pegawai!

    peg_max = ws_peg.max_row if ws_peg.max_row > 1 else 6587
    wb.defined_names['Daftar_Nama'] = DefinedName('Daftar_Nama', attr_text=f'Pegawai!$A$2:$A${peg_max}')
    wb.defined_names['Daftar_Bidang'] = DefinedName('Daftar_Bidang', attr_text='Pegawai!$E$2:$E$7')

    # 2. Sheet DATA
    ws=wb.create_sheet('DATA');ws.sheet_view.showGridLines=False
    ws['A1']='DATA MASTER PENUGASAN — MPP BPKP';ws.merge_cells('A1:D1')
    ws['A2']='Kuning = isian pengguna; Abu-abu = formula otomatis. Contoh isian nyata tersedia langsung pada Kolom D.';ws.merge_cells('A2:D2');ws.row_dimensions[2].height=32
    
    for addr,value in [('A3','TIM PENGAWASAN'),('B3','Nama Pegawai'),('C3','NIP (Otomatis)'),('D3','Petunjuk & Contoh Isian')]:ws[addr]=value
    
    # Rows 4 to 26: Tim Pengawasan (PM, PT, KT, AT 1..20)
    for r,role in enumerate(BASE_ROLES,4):
        ws.cell(r,1,role);editable(ws.cell(r,2))
        calculated(ws.cell(r,3),f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r}, Pegawai!$A:$B, 2, FALSE), ""))',explanation='NIP terisi otomatis saat Nama dipilih dari dropdown.')
        ws.cell(r,3).number_format='@'
        if r<=6:
            guide_cell(ws.cell(r,4),'Pilih nama dari dropdown atau ketik nama (NIP terisi otomatis). Wajib diisi.')
        else:
            guide_cell(ws.cell(r,4),f'Anggota Tim {r-6}: Pilih nama dari dropdown (baris anggota yang tidak digunakan cukup dikosongkan).')
    
    # Row 27: Kepala Perwakilan
    ws['A27']='Kepala Perwakilan';editable(ws['B27'])
    calculated(ws['C27'],'=IF(B27="","",IFERROR(VLOOKUP(B27, Pegawai!$A:$B, 2, FALSE), ""))')
    ws['C27'].number_format='@'
    guide_cell(ws['D27'],'Pilih nama Kepala Perwakilan penandatangan Surat Tugas (NIP terisi otomatis).')

    # Data Validation for Names (B4:B27)
    dv_nama = DataValidation(type='list', formula1='Daftar_Nama', allow_blank=True, showInputMessage=True, showErrorMessage=True, errorTitle='Peringatan Nama Pegawai', error='Nama tidak terdaftar dalam database pegawai. Tetap gunakan?', promptTitle='Pilih Nama Pegawai', prompt='Ketik atau pilih nama pegawai dari daftar dropdown.', errorStyle='warning')
    ws.add_data_validation(dv_nama);dv_nama.add('B4:B27')

    # Data Validation for Bidang (B29)
    dv_bidang = DataValidation(type='list', formula1='Daftar_Bidang', allow_blank=True, showInputMessage=True, showErrorMessage=True, errorTitle='Bidang Tidak Valid', error='Silakan pilih salah satu bidang dari daftar yang tersedia.', promptTitle='Pilih Bidang', prompt='Pilih bidang penugasan dari daftar dropdown.', errorStyle='warning')
    ws.add_data_validation(dv_bidang);dv_bidang.add('B29')

    contoh_isian = {
        29: 'Pilihan bidang: Umum, IPP, APD, AN, Investigasi, P3A (Gunakan dropdown).',
        30: 'Contoh: Evaluasi atas Tata Kelola Program Ketahanan Pangan pada Pemerintah Kabupaten Nabire',
        31: 'Contoh: Menilai efektivitas dan akuntabilitas penyaluran bantuan cadangan pangan pemerintah daerah',
        32: 'Contoh: KP-01/PW36/2/2026 (Opsional, boleh kosong saat konsep)',
        33: 'Contoh: Dinas Ketahanan Pangan Kabupaten Nabire',
        34: 'Contoh: Triwulan III Tahun 2026 (atau Tahun Anggaran 2026)',
        35: 'Contoh: R-01/PW36/2/2026 (Opsional / sesuai kondisi)',
        36: 'Pilihan tingkat risiko: Rendah, Sedang, Tinggi (Gunakan dropdown, default: Sedang)',
        37: 'Contoh: Jl. Merdeka No. 1, Nabire, Papua Tengah (Telp. 0984-21001)',
        39: 'Contoh: 02/11/2026 (Format tanggal Surat Tugas)',
        40: 'Otomatis: sama dengan Tanggal Konsep ST',
        41: 'Contoh: 1 (Jumlah hari kerja kalender fase persiapan)',
        42: 'Contoh: 3 (Jumlah hari kerja kalender fase pelaksanaan)',
        43: 'Contoh: 1 (Jumlah hari kerja kalender fase penyelesaian)',
        44: 'Otomatis: dihitung 2 hari kerja setelah tanggal ST',
        45: 'Otomatis: rentang hari persiapan',
        46: 'Otomatis: awal pelaksanaan',
        47: 'Otomatis: akhir pelaksanaan (Sabtu dan Minggu dilewati otomatis)',
        48: 'Otomatis: awal penyelesaian',
        49: 'Otomatis: akhir penyelesaian',
        50: 'Otomatis: direncanakan selesai penugasan ST',
        51: 'Otomatis: Rencana Penerbitan Laporan (RPL)',
        53: 'Contoh: 6.5 (Standar jam pengawasan per hari orang BPKP)',
        54: 'Otomatis: menghitung jumlah anggota tim yang terisi nama',
        56: 'Pilihan: PIA atau PENGAWASAN (Gunakan dropdown)',
        57: 'Pilihan: EVALUASI, AUDIT, PIA, LAINNYA (Gunakan dropdown, default: EVALUASI)',
        58: 'Contoh: Pertanian dan Ketahanan Pangan',
        59: 'Contoh: Penguatan Ketahanan Pangan Daerah',
        60: 'Contoh: Distribusi Cadangan Beras Pemerintah',
        61: 'Contoh: Memperoleh keyakinan memadai atas ketepatan sasaran dan kepatuhan penyaluran bantuan cadangan pangan',
        62: 'Contoh: Penyaluran bantuan pangan periode Januari s.d. September 2026 di Kabupaten Nabire',
        63: 'Contoh: Bupati Nabire',
        64: 'Contoh: DIPA Perwakilan BPKP Provinsi Papua Tengah Tahun Anggaran 2026',
        65: 'Contoh: Peraturan BPKP Nomor 2 Tahun 2023 tentang Manajemen Penugasan Pengawasan di Lingkungan BPKP',
        66: 'Contoh: Pedoman Pengawasan Ketahanan Pangan Kedeputian Bidang Perekonomian BPKP',
        67: 'Contoh: Kepala Dinas Ketahanan Pangan Kabupaten Nabire',
        68: 'Contoh: Nabire',
        69: 'Contoh: Penyampaian Surat Tugas Pengawasan (Opsional, kosong = Nama Penugasan)',
        70: 'Contoh: PR.01.01/PW36/2/2026 (Opsional, boleh kosong saat konsep)',
        71: 'Otomatis: sama dengan Tanggal ST',
        72: 'Contoh: ST-01/PW36/2/2026 (Opsional, boleh kosong saat konsep)',
        73: 'Contoh: Kepala Perwakilan (Default)',
        74: 'Otomatis: sama dengan Tanggal Pengesahan',
        75: 'Otomatis: tanggal mulai periode ST',
        76: 'Otomatis: total hari kerja ST (Sabtu dan Minggu dilewati otomatis)',
        77: 'Contoh: RMP-01 (Opsional, jika ada menurut PKPT)',
        78: 'Contoh: RPL-01 (Opsional, jika ada menurut PKPT)',
        79: 'Contoh: RPL-NON-01 (Opsional, jika ada)',
        80: 'Otomatis: target selesai konsep laporan',
    }

    for r,(key,label,value,required) in DATA.items():
        ws.cell(r,1,label);editable(ws.cell(r,2),value if value!='' else None)
        guide_cell(ws.cell(r,4), contoh_isian.get(r, 'Wajib diisi' if required else 'Opsional / sesuai kondisi'))
        ws.cell(r,5,key)

    for r,label in [(28,'DETAIL PENUGASAN'),(38,'JADWAL PENUGASAN'),(52,'PARAMETER WAKTU'),(55,'KELENGKAPAN SURAT & ADMINISTRASI'),(81,'SUMBER DAN SKEMA')]:ws.cell(r,1,label)
    ws['A82']='Versi skema';ws['B82']=VERSION
    ws['A83']='Workbook sumber';ws['B83']='D:/autoST/KM4_KM5_KM6.xlsx — DATA (dibaca, tidak diubah)'

    formulas={
        40:'=IF(B39="","",B39)',
        44:'=IF(B39="","",WORKDAY(B39,2))',
        45:'=IF(OR(B44="",B41=""),"",WORKDAY(B44,B41-1))',
        46:'=IF(B45="","",WORKDAY(B45,1))',
        47:'=IF(OR(B46="",B42=""),"",WORKDAY(B46,B42-1))',
        48:'=IF(B47="","",WORKDAY(B47,1))',
        49:'=IF(OR(B48="",B43=""),"",WORKDAY(B48,B43-1))',
        50:'=IF(B49="","",B49)',
        51:'=IF(B50="","",B50)',
        54:'=COUNTA(B7:B26)',
        71:'=IF(B39="","",B39)',
        74:'=IF(B40="","",B40)',
        75:'=IF(B44="","",B44)',
        76:'=IF(OR(B75="",B50=""),"",NETWORKDAYS(B75,B50))',
        80:'=IF(B50="","",B50)',
    }
    for r,f in formulas.items():calculated(ws.cell(r,2),f,'Nilai otomatis; Sabtu/Minggu dikecualikan.')
    ws['B53'].comment=Comment('6,5 jam per hari orang berasal dari parameter workbook sumber. Sesuaikan jika dasar penugasan berbeda.','MPP BPKP')
    
    for r in [39,40,44,45,46,47,48,49,50,51,71,74,75,77,78,79,80]:ws.cell(r,2).number_format='dd mmmm yyyy'
    for r in [30,31,33,61,62,64,65,66]:ws.row_dimensions[r].height=52
    for col,width in [('A',38),('B',52),('C',26),('D',82)]:ws.column_dimensions[col].width=width
    ws.column_dimensions['E'].hidden=True;ws.freeze_panes='B4'

    dropdown(ws,['B36'],['Rendah','Sedang','Tinggi'])
    dropdown(ws,['B56'],['PIA','PENGAWASAN'])
    dropdown(ws,['B57'],['EVALUASI','AUDIT','PIA','LAINNYA'])

    for name,headers in TABLES.items():
        if name not in ['IHP','TEMBUSAN']:continue
        s=wb.create_sheet(name)
        note='Ketik butir Insilwas per baris (boleh diawali ID seperti IHP-01, nomor 1., atau langsung uraian; skrip otomatis mengekstrak ID).' if name=='IHP' else 'Isi baris mulai 5; baris kosong tidak diproses. Tambahkan baris bila diperlukan.'
        base(s,name.replace('_',' '),note,headers)
        n=30
        for row in range(5,n+5):
            for col in range(1,len(headers)+1):editable(s.cell(row,col))
            s.row_dimensions[row].height=45 if name=='IHP' else (28 if name=='TEMBUSAN' else 65)
        for col in range(1,len(headers)+1):s.column_dimensions[chr(64+col)].width=110 if name=='IHP' else (55 if name!='PROGRAM_KERJA' else [16,48,95,35,14,14,20][col-1])
        s.auto_filter.ref=f'A4:{chr(64+len(headers))}{n+4}'

    s=wb.create_sheet('JAM_TIM')
    base(s,'ALOKASI JAM TIAP PERSONEL','Awal: partisipasi penuh × hari fase × jam/hari. Tinjau PM/PT/anggota dan ubah jam fase di E:G bila berbeda. Anggaran, bukan realisasi.',['No.','Peran','Nama','NIP','Persiapan (jam)','Pelaksanaan (jam)','Penyelesaian (jam)','Total (jam)'])
    for r in range(5,len(BASE_ROLES)+5):
        idx=r-4
        role=BASE_ROLES[idx-1]
        refname=f'DATA!B{idx+3}'
        refnip=f'DATA!C{idx+3}'
        s.cell(r,1,idx);s.cell(r,2,role)
        calculated(s.cell(r,3),f'=IF({refname}="","",{refname})')
        calculated(s.cell(r,4),f'=IF({refname}="","",{refnip}&"")');s.cell(r,4).number_format='@'
        for col,phase in [(5,41),(6,42),(7,43)]:calculated(s.cell(r,col),f'=IF(C{r}="","",DATA!$B${phase}*DATA!$B$53)','Asumsi awal partisipasi penuh. Ganti jam sesuai rencana orang ini.')
        calculated(s.cell(r,8),f'=IF(C{r}="","",SUM(E{r}:G{r}))')
        for col in range(5,9):s.cell(r,col).number_format='0.00'
    for col,width in [('A',8),('B',30),('C',50),('D',26),('E',22),('F',22),('G',22),('H',22)]:s.column_dimensions[col].width=width
    for s in wb:
        for row in s:
            for c in row:
                if c.font.name!='Arial':c.font=Font(name='Arial',size=11,bold=c.row in [1,3,4],color='000000')
                if not c.alignment.wrap_text:c.alignment=Alignment(vertical='top',wrap_text=True)
        s.sheet_properties.outlinePr.summaryRight=False
    for r in [1,3,28,38,52,55,81]:
        for c in ws[r][:4]:c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name='Arial',size=11,bold=True,color=WHITE)
    ws.print_area='A1:D83';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    wb.active = ws
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists() and not overwrite:raise FileExistsError(f'Tidak menimpa workbook yang sudah ada: {path}')
    wb.save(path);print(path)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('output',type=Path);args=p.parse_args();create(args.output)
