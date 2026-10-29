"""Fakta workbook terpisah dari rancangan analitis yang disusun agent."""
from datetime import date, datetime, timedelta
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import tempfile
from zipfile import ZipFile
from lxml import etree as E
import openpyxl
from schema import ROOT, VERSION, DATA, NARASI, BASE_ROLES, PIA_REQUIRED
from office_io import recalculate_xlsx

MONTHS=['','Januari','Februari','Maret','April','Mei','Juni','Juli','Agustus','September','Oktober','November','Desember']
class InputError(Exception):
    def __init__(self,issues):self.issues=issues;super().__init__('; '.join(issues))
def nonempty(v):return v is not None and str(v).strip()!=''
def fmt(v):
    if v is None:return ''
    if isinstance(v,(datetime,date)):return f'{v.day} {MONTHS[v.month]} {v.year}'
    if isinstance(v,float):return f'{v:.3f}'.rstrip('0').rstrip('.').replace('.',',')
    return str(v).strip()
def lines(v):return [s.strip() for s in str(v or '').splitlines() if s.strip()]
def words(n):
    n=int(n);a=['nol','satu','dua','tiga','empat','lima','enam','tujuh','delapan','sembilan','sepuluh','sebelas']
    if n<12:return a[n]
    if n<20:return words(n-10)+' belas'
    if n<100:return words(n//10)+' puluh'+(' '+words(n%10) if n%10 else '')
    if n<200:return 'seratus'+(' '+words(n-100) if n>100 else '')
    if n<1000:return words(n//100)+' ratus'+(' '+words(n%100) if n%100 else '')
    if n<2000:return 'seribu'+(' '+words(n-1000) if n>1000 else '')
    if n<1000000:return words(n//1000)+' ribu'+(' '+words(n%1000) if n%1000 else '')
    return str(n)
def number(v,label,issues,positive=False):
    try:
        if isinstance(v,bool):raise ValueError()
        n=float(str(v).replace(',','.'))
        if not math.isfinite(n) or n<0 or (positive and n<=0):raise ValueError()
        return n
    except (ValueError,TypeError):issues.append(label+': angka '+('lebih dari nol' if positive else 'nol atau positif')+' diperlukan.');return 0.
def as_date(v,label,issues):
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    if isinstance(v,str):
        for pattern in ['%Y-%m-%d','%d/%m/%Y','%d-%m-%Y']:
            try:return datetime.strptime(v.strip(),pattern).date()
            except ValueError:pass
    issues.append(label+': isi tanggal Excel atau YYYY-MM-DD.');return None
def table_rows(ws,n):
    return [(r,[ws.cell(r,c).value for c in range(1,n+1)]) for r in range(5,ws.max_row+1) if any(nonempty(ws.cell(r,c).value) for c in range(1,n+1))]
def source_path(value,workbook,label,issues,required=False,office_only=True):
    if not nonempty(value):
        if required:issues.append(label+': lokasi berkas wajib diisi.')
        return None
    p=Path(str(value).strip().strip('"')).expanduser()
    if not p.is_absolute():p=workbook.parent/p
    if not p.is_file():issues.append(f'{label}: berkas tidak ditemukan: {p}');return None
    suffix=p.suffix.lower()
    if office_only and suffix not in ['.docx','.pdf']:issues.append(label+': gunakan DOCX/PDF.');return None
    try:
        if suffix=='.docx':
            with ZipFile(p) as z:E.fromstring(z.read('word/document.xml'))
        elif suffix=='.pdf':
            import fitz
            with fitz.open(p) as d:
                if not len(d):raise ValueError('tanpa halaman')
    except Exception as e:issues.append(f'{label}: berkas tidak dapat dibaca ({e}).');return None
    return p.resolve()

def read_model(workbook,recalc=True,content=None,inspect_only=False):
    workbook=Path(workbook).resolve();issues=[]
    with tempfile.TemporaryDirectory(prefix='mpp-input-') as tmp:
        snapshot=Path(tmp)/'isian.xlsx';shutil.copy2(workbook,snapshot)
        if recalc:recalculate_xlsx(snapshot)
        wb=openpyxl.load_workbook(snapshot,data_only=True)
        required=['DATA','JAM_TIM','IHP','TEMBUSAN']
        missing=[s for s in required if s not in wb.sheetnames]
        if missing:raise InputError(['Gunakan Isian_MPP_BPKP.xlsx baru. Tab belum ada: '+', '.join(missing)])
        for s in wb:
            for row in s:
                for c in row:
                    if c.data_type=='e':issues.append(f'{s.title}!{c.coordinate}: formula menghasilkan {c.value}.')
        if issues:raise InputError(issues)
        key_to_row={}
        for r in range(1,wb['DATA'].max_row+1):
            k=wb['DATA'].cell(r,5).value
            if k and isinstance(k,str) and k.strip():key_to_row[k.strip()]=r
        data={}
        for def_r,(key,_,_,_) in DATA.items():
            actual_r=key_to_row.get(key,def_r)
            data[key]=wb['DATA'].cell(actual_r,2).value
        version_val=None
        for r in range(wb['DATA'].max_row,1,-1):
            if wb['DATA'].cell(r,1).value=='Versi skema':version_val=wb['DATA'].cell(r,2).value;break
        if version_val and fmt(version_val)!=VERSION:issues.append('Versi skema tidak dikenali.')
        originals={}
        try:
            originals=json.loads((ROOT/'assets'/'profil_unit.json').read_text(encoding='utf-8'))
        except Exception:pass
        if not nonempty(data.get('UNIT_KERJA')):data['UNIT_KERJA']=originals.get('unit','Perwakilan BPKP Provinsi Papua Tengah')
        if not nonempty(data.get('KOP_BARIS1')):data['KOP_BARIS1']=originals.get('nama_kop_baris1','BADAN PENGAWASAN KEUANGAN DAN PEMBANGUNAN')
        if not nonempty(data.get('KOP_BARIS2')):data['KOP_BARIS2']=originals.get('nama_kop_baris2','PERWAKILAN PROVINSI PAPUA TENGAH')
        if not nonempty(data.get('ALAMAT_UNIT')):data['ALAMAT_UNIT']=originals.get('alamat','Jalan Sam Ratulangi, Kelurahan Oyehe, Distrik Nabire')
        if not nonempty(data.get('WILAYAH_KODEPOS')):data['WILAYAH_KODEPOS']=originals.get('wilayah_kodepos','Kabupaten Nabire, Provinsi Papua Tengah, Kode Pos: 98816')
        if not nonempty(data.get('EMAIL_UNIT')):data['EMAIL_UNIT']=originals.get('email','papua.tengah@bpkp.go.id')
        if not nonempty(data.get('WEBSITE_UNIT')):data['WEBSITE_UNIT']=originals.get('website','www.bpkp.go.id/papua.tengah')
        if not nonempty(data.get('KOTA_UNIT')):data['KOTA_UNIT']=originals.get('kota','Nabire')
        if not nonempty(data.get('JALUR_ST')):data['JALUR_ST']='MANUAL'
        if not nonempty(data.get('JABATAN_KAPER')):data['JABATAN_KAPER']='Kepala Perwakilan'
        if not nonempty(data.get('TINGKAT_RISIKO')):data['TINGKAT_RISIKO']='Sedang'
        for def_r,(key,label,_,required_value) in DATA.items():
            actual_r=key_to_row.get(key,def_r)
            if required_value and not nonempty(data.get(key)):issues.append(f'DATA!B{actual_r}: {label} belum diisi.')
        jenis=fmt(data['JENIS_PENUGASAN']).upper();jalur=fmt(data['JALUR_ST']).upper()
        if jenis not in ['PIA','PENGAWASAN']:issues.append('DATA: jenis penugasan harus PIA atau PENGAWASAN.')
        if jalur not in ['MANUAL','SIMA-NG']:issues.append('DATA: jalur ST harus MANUAL atau SIMA-NG.')
        if not nonempty(data.get('BENTUK_PENGAWASAN')):
            data['BENTUK_PENGAWASAN']='PIA' if jenis=='PIA' else 'EVALUASI'
        if not nonempty(data.get('STATUS_CEK_DPP')):data['STATUS_CEK_DPP']='TERSEDIA'
        if not nonempty(data.get('STATUS_CEK_RENCANA')):data['STATUS_CEK_RENCANA']='TERSEDIA'
        if not nonempty(data.get('STATUS_CEK_KM6')):data['STATUS_CEK_KM6']='TERSEDIA'
        if not nonempty(data.get('STATUS_CEK_PIA')):
            data['STATUS_CEK_PIA']='N/A' if jenis=='PIA' else 'TERSEDIA'
        if fmt(data['BENTUK_PENGAWASAN']).upper() not in ['AUDIT','EVALUASI','PIA','LAINNYA']:issues.append('DATA: pilih AUDIT, EVALUASI, PIA, atau LAINNYA.')
        if fmt(data['TINGKAT_RISIKO']) not in ['Rendah','Sedang','Tinggi']:issues.append('DATA: tingkat risiko harus Rendah, Sedang, atau Tinggi.')
        kaper_row=27
        for r in range(4,35):
            if 'Kepala Perwakilan' in str(wb['DATA'].cell(r,1).value or ''):
                kaper_row=r;break
        data['NAMA_KAPER']=wb['DATA'].cell(kaper_row,2).value
        if not nonempty(data['NAMA_KAPER']):issues.append(f'DATA!B{kaper_row}: nama Kepala Perwakilan belum diisi.')
        team=[]
        def person(name,nip,role,label,gol='-'):
            if not nonempty(name):return
            clean_nip=re.sub(r'\s+','',str(nip or '').strip())
            if not nonempty(nip) or len(clean_nip)!=18 or not clean_nip.isdigit():issues.append(label+': NIP harus 18 digit angka.')
            team.append({'name':fmt(name),'nip':fmt(nip).strip(),'clean_nip':clean_nip,'gol':fmt(gol or '-'),'role':role})
        for r,role in [(4,'Wakil Penanggung Jawab'),(5,'Pengendali Teknis'),(6,'Ketua Tim')]:
            if not nonempty(wb['DATA'].cell(r,2).value):issues.append(f'DATA!B{r}: peran {role} wajib diisi.')
            person(wb['DATA'].cell(r,2).value,wb['DATA'].cell(r,3).value,role,f'DATA baris {r}')
        for r in range(7,kaper_row):
            lbl=str(wb['DATA'].cell(r,1).value or '')
            if 'Anggota' in lbl or 'AT' in lbl:
                name=wb['DATA'].cell(r,2).value;nip=wb['DATA'].cell(r,3).value
                if nonempty(name):person(name,nip,'Anggota Tim',f'DATA baris {r}')
        if 'TIM_TAMBAHAN' in wb.sheetnames:
            for r,row in table_rows(wb['TIM_TAMBAHAN'],3):
                if nonempty(row[0]):person(row[0],row[1] if len(row)>1 else '','Anggota Tim',f'TIM_TAMBAHAN baris {r}',row[2] if len(row)>2 else '-')
        if len({p['clean_nip'] for p in team})!=len(team):issues.append('Ada NIP tim duplikat.')
        if len({p['name'] for p in team})!=len(team):issues.append('Nama tim harus unik agar pembagian tugas tidak ambigu.')
        allocations={}
        for r in range(5,wb['JAM_TIM'].max_row+1):
            if nonempty(wb['JAM_TIM'].cell(r,3).value):
                raw_nip=fmt(wb['JAM_TIM'].cell(r,4).value)
                clean_k=re.sub(r'\s+','',raw_nip)
                if clean_k:
                    if clean_k in allocations:issues.append(f'JAM_TIM baris {r}: NIP duplikat.')
                    allocations[clean_k]=[number(wb['JAM_TIM'].cell(r,c).value,f'JAM_TIM!{openpyxl.utils.get_column_letter(c)}{r}',issues) for c in [5,6,7]]
        for p in team:
            if p['clean_nip'] not in allocations:issues.append('JAM_TIM: alokasi belum tersedia untuk '+p['name'])
            p['hours']=allocations.get(p['clean_nip'],[0.,0.,0.])
        if set(allocations)-{p['clean_nip'] for p in team}:issues.append('JAM_TIM memuat personel yang tidak ada dalam tim.')
        r_jam=key_to_row.get('JAM_PER_HARI',47)
        jam_day=number(data['JAM_PER_HARI'],f'DATA!B{r_jam}',issues,True)
        dates={}
        for key in ['TANGGAL_ST','TANGGAL_PENGESAHAN','PERS_MULAI','PERS_SELESAI','LAKS_MULAI','LAKS_SELESAI','SELESAI_MULAI','SELESAI_AKHIR','TANGGAL_MULAI','TANGGAL_SELESAI','RPL','TANGGAL_PAKTA','TANGGAL_SURAT_PENGANTAR','TGL_KONSEP_LAPORAN']:
            if nonempty(data.get(key)):dates[key]=as_date(data[key],key,issues)
        holidays=set()
        if 'LIBUR' in wb.sheetnames:
            for r,row in table_rows(wb['LIBUR'],2):
                d=as_date(row[0],f'LIBUR!A{r}',issues)
                if d:holidays.add(d)
        if len(dates)==14 and all(dates.values()):
            order=['TANGGAL_ST','PERS_MULAI','PERS_SELESAI','LAKS_MULAI','LAKS_SELESAI','SELESAI_MULAI','SELESAI_AKHIR']
            if any(dates[a]>dates[b] for a,b in zip(order,order[1:])):issues.append('Urutan fase tanggal tidak konsisten.')
            if dates['TANGGAL_MULAI']>dates['PERS_MULAI'] or dates['TANGGAL_SELESAI']<dates['SELESAI_AKHIR']:issues.append('Periode ST harus mencakup seluruh fase.')
            def workdays(a,b):return sum(1 for i in range(max(0,(b-a).days+1)) if (a+timedelta(days=i)).weekday()<5 and a+timedelta(days=i) not in holidays)
            days=workdays(dates['TANGGAL_MULAI'],dates['TANGGAL_SELESAI'])
            r_hari=key_to_row.get('JUMLAH_HARI',77)
            if number(data['JUMLAH_HARI'],f'DATA!B{r_hari}',issues)!=days:issues.append(f'DATA!B{r_hari} tidak cocok dengan hari kerja ST.')
            for start,end,duration in [('PERS_MULAI','PERS_SELESAI','HARI_PERSIAPAN'),('LAKS_MULAI','LAKS_SELESAI','HARI_PELAKSANAAN'),('SELESAI_MULAI','SELESAI_AKHIR','HARI_PENYELESAIAN')]:
                if number(data[duration],duration,issues,True)!=workdays(dates[start],dates[end]):issues.append(duration+' tidak sama dengan hari kerja rentang fasenya.')
            if dates['TGL_KONSEP_LAPORAN']>dates['RPL']:issues.append('Target konsep laporan melewati RPL.')
            if len(team)>=3 and jam_day and not math.isclose(sum(team[2]['hours'])/jam_day,days,abs_tol=.001):issues.append('Hari orang KT pada KM 4 harus sama dengan hari kerja ST.')
        ihp=[]
        for r in range(5,wb['IHP'].max_row+1):
            val_a=wb['IHP'].cell(r,1).value;val_b=wb['IHP'].cell(r,2).value
            if nonempty(val_a) and nonempty(val_b):
                ihp.append({'id':fmt(val_a),'information':fmt(val_b),'source':fmt(wb['IHP'].cell(r,3).value or '')})
            elif nonempty(val_a):
                entry_text=str(val_a).strip()
                m_prefix=re.match(r'^(?:\[?(IHP[-._\s]?\d+)\]?|[(\[]?(\d+)[)\]]?)\s*[-–—:.)]\s*(.*)$',entry_text,re.IGNORECASE)
                if not m_prefix:m_prefix=re.match(r'^(IHP[-._\s]?\d+)\s+(.*)$',entry_text,re.IGNORECASE)
                if m_prefix:
                    raw_id=m_prefix.group(1) if m_prefix.group(1) else m_prefix.group(2)
                    clean_text=m_prefix.groups()[-1].strip() if m_prefix.groups()[-1] else entry_text
                    num_m=re.search(r'\d+',raw_id)
                    ihp_id=f'IHP-{int(num_m.group()):02d}' if num_m else f'IHP-{len(ihp)+1:02d}'
                    ihp.append({'id':ihp_id,'information':clean_text or entry_text,'source':''})
                else:
                    ihp.append({'id':f'IHP-{len(ihp)+1:02d}','information':entry_text,'source':''})
        files={};sources={}
        files={};sources={}
        if jalur=='SIMA-NG' and nonempty(data.get('FILE_ST_SIMA')):
            r_st_sima=key_to_row.get('FILE_ST_SIMA',0)
            files[6]=source_path(data['FILE_ST_SIMA'],workbook,f'DATA!B{r_st_sima}',issues,True)
            for i,key in [(7,'FILE_KM4_SIMA'),(8,'FILE_KM5_SIMA')]:
                if nonempty(data.get(key)):files[i]=source_path(data[key],workbook,f'DATA!B{key_to_row.get(key,0)}',issues)
        if jenis=='PENGAWASAN':
            sumber_pia=fmt(data.get('SUMBER_PIA') or 'ISIAN').upper()
            if sumber_pia=='FILE' and nonempty(data.get('FILE_LAPORAN_PIA')):
                r_pia=key_to_row.get('FILE_LAPORAN_PIA',0)
                files[9]=source_path(data['FILE_LAPORAN_PIA'],workbook,f'DATA!B{r_pia}',issues,True)
        for key in ['FILE_KAP','FILE_PEDOMAN']:
            if nonempty(data.get(key)):sources[key]=source_path(data[key],workbook,f'DATA!B{key_to_row.get(key,0)}',issues,office_only=False)
        if not ihp and not nonempty(data.get('FILE_KAP')):issues.append('Isi Insilwas pada tab IHP.')
        if data.get('UNIT_KERJA') and data.get('KOP_BARIS2') and fmt(data['UNIT_KERJA'])!='Perwakilan BPKP Provinsi Papua Tengah' and 'PAPUA TENGAH' in fmt(data['KOP_BARIS2']).upper():issues.append('Unit diubah tetapi kop masih Papua Tengah.')
        for key in ['STATUS_CEK_PIA','STATUS_CEK_DPP','STATUS_CEK_RENCANA','STATUS_CEK_KM6']:
            if nonempty(data.get(key)) and fmt(data[key]).upper() not in ['TERSEDIA','BELUM','N/A']:issues.append(key+': gunakan TERSEDIA, BELUM, atau kosong.')
        if issues:raise InputError(issues)
        model={'data':data,'team':team,'tables':{'IHP':[[x['information'],x['source']] for x in ihp],'TEMBUSAN':[row for _,row in table_rows(wb['TEMBUSAN'],1)]},'ihp':ihp,'sources':sources,'pia':{},'files':files,'jenis':jenis,'jalur':jalur,'stages':[sum(p['hours'][i] for p in team) for i in range(3)],'jam_day':jam_day,'workbook':workbook,'holidays':holidays,'input_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()}
    if inspect_only:return model
    if not content:raise InputError(['Rancangan analitis belum disusun oleh skill. Jalankan --inspect, susun risiko dan tugas dari sumber, lalu gunakan --content rancangan.json. Pengguna tidak perlu mengisi risiko/program kerja pada XLSX.'])
    plan=json.loads(Path(content).read_text(encoding='utf-8-sig')) if not isinstance(content,dict) else content
    return compile_plan(model,plan)

def compile_plan(m,plan):
    issues=[]
    if plan.get('schema_version')!=VERSION:issues.append('Rancangan: schema_version harus '+VERSION)
    if plan.get('input_sha256')!=m['input_sha256']:issues.append('Rancangan berasal dari versi input berbeda; susun ulang terhadap workbook terbaru.')
    ihp=plan.get('ihp',m['ihp']);ids=[x.get('id','') for x in ihp]
    if not ihp or len(set(ids))!=len(ids) or any(not nonempty(x.get('id')) or not nonempty(x.get('information')) for x in ihp):issues.append('Rancangan: Insilwas harus memiliki ID unik dan uraian.')
    existing={x['id']:x['information'] for x in m['ihp']}
    for x in ihp:
        if x.get('id') in existing and x.get('information')!=existing[x['id']]:issues.append('Insilwas workbook tidak boleh diubah diam-diam: '+x['id'])
    if set(existing)-set(ids):issues.append('Ada Insilwas workbook yang tidak dimasukkan ke rancangan.')
    narratives=plan.get('narratives',{})
    for key,(label,required) in NARASI.items():
        if required and not nonempty(narratives.get(key)):issues.append('Rancangan narasi belum lengkap: '+key)
    processes=plan.get('processes',[])
    if not processes or any(len(row)!=5 or not all(nonempty(v) for v in row) for row in processes):issues.append('Rancangan proses bisnis: minimal satu baris dengan tahap/tujuan/aktivitas/input/output.')
    risks=plan.get('client_risks',[]);risk_ids=[];covered=set();risk_rows=[]
    for x in risks:
        rid=x.get('id','');risk_ids.append(rid);links=x.get('ihp_ids',[])
        if not links or set(links)-set(ids):issues.append(f'Risiko klien {rid}: hubungan Insilwas tidak sah/kosong.')
        covered.update(links)
        for k in ['id','event','cause','impact']:
            if not nonempty(x.get(k)):issues.append(f'Risiko klien {rid}: {k} belum disusun.')
        risk_rows.append([x.get('event',''),'ID: '+rid+'; Insilwas: '+', '.join(links)+'\nPenyebab: '+x.get('cause','')+ ('\nDasar/indikasi: '+x['basis'] if x.get('basis') else ''),x.get('impact','')])
    if not risks or len(set(risk_ids))!=len(risk_ids) or any(not x for x in risk_ids):issues.append('Risiko klien: ID unik dan minimal satu risiko diperlukan.')
    if set(ids)-covered:issues.append('Insilwas belum memiliki risiko klien: '+', '.join(sorted(set(ids)-covered)))
    people={p['name']:p for p in m['team']};assignment=[]
    for x in plan.get('assignment_risks',[]):
        for k in ['event','cause','impact','controls','owner','when']:
            if not nonempty(x.get(k)):issues.append('Risiko penugasan: '+k+' belum disusun.')
        if x.get('owner') not in people:issues.append('Risiko penugasan: pemilik pengendalian harus nama anggota tim.')
        assignment.append([x.get('event','')+'\nPenyebab: '+x.get('cause',''),x.get('impact',''),x.get('controls','')+'\nPemilik: '+x.get('owner','')+'\nWaktu: '+x.get('when','')])
    if not assignment:issues.append('Minimal satu risiko penugasan beserta pengendaliannya diperlukan.')
    objectives=plan.get('objectives',[]);objective_map={};covered_by_objectives=set()
    expected={'AUDIT':'TAO','EVALUASI':'TEO'}.get(fmt(m['data']['BENTUK_PENGAWASAN']).upper())
    for objective in objectives:
        oid=objective.get('id','');links=objective.get('risk_ids',[])
        if not oid or oid in objective_map:issues.append('TAO/TEO: ID harus unik dan terisi.')
        objective_map[oid]=objective
        if not links or set(links)-set(risk_ids):issues.append('TAO/TEO '+oid+': risiko klien belum dipetakan dengan benar.')
        covered_by_objectives.update(links)
        if expected and objective.get('type')!=expected:issues.append(oid+': jenis tujuan harus '+expected+' sesuai bentuk penugasannya.')
        if not nonempty(objective.get('text')):issues.append('TAO/TEO '+oid+': rumusan tujuan belum ada.')
        source=objective.get('source',{})
        if source.get('kind') not in ['PEDOMAN_PUSAT','USULAN_TENTATIF']:issues.append('TAO/TEO '+oid+': tandai sumber pedoman pusat atau usulan tentatif.')
        if source.get('kind')=='PEDOMAN_PUSAT':
            if not nonempty(source.get('path')) or not nonempty(source.get('location')):issues.append('TAO/TEO '+oid+': lokasi pedoman pusat dan halaman/bagiannya wajib dicatat.')
            elif not Path(source['path']).is_file():issues.append('TAO/TEO '+oid+': berkas pedoman pusat tidak ditemukan.')
        elif not nonempty(source.get('basis')):issues.append('TAO/TEO '+oid+': usulan tentatif harus memiliki dasar dan tidak boleh diklaim sebagai rumusan pusat.')
    if not objectives:issues.append('Rancangan belum memuat TAO/TEO atau tujuan tentatif sesuai jenis penugasan.')
    if set(risk_ids)-covered_by_objectives:issues.append('Risiko klien belum dipetakan ke TAO/TEO: '+', '.join(sorted(set(risk_ids)-covered_by_objectives)))
    tasks=[];covered_objectives=set();step_ids=[]
    for task in plan.get('tasks',[]):
        links=task.get('objective_ids',[]);tid=task.get('id','')
        if not links or set(links)-set(objective_map):issues.append('KM 6 '+tid+': tautan TAO/TEO kosong/tidak sah.')
        covered_objectives.update(links)
        risk_links=sorted({rid for oid in links if oid in objective_map for rid in objective_map[oid].get('risk_ids',[])})
        if not nonempty(task.get('objective')):issues.append('KM 6 '+tid+': tujuan pengujian belum disusun.')
        if not task.get('steps'):issues.append('KM 6 '+tid+': belum ada langkah rinci.')
        for i,step in enumerate(task.get('steps',[]),1):
            sid=step.get('id',f'{tid}.{i:02}');step_ids.append(sid)
            names=step.get('assignees',[])
            if not names or not isinstance(names,list) or any(n not in people for n in names):issues.append('KM 6 '+sid+': setiap task harus ditugaskan kepada nama personel yang ada di tim, bukan label Tim/KT/AT.')
            if len(set(names))!=len(names):issues.append('KM 6 '+sid+': nama pelaksana duplikat.')
            if not nonempty(step.get('text')) or not nonempty(step.get('reference')):issues.append('KM 6 '+sid+': langkah kerja dan referensi kertas kerja wajib disusun.')
            weight=number(step.get('weight',1),'Bobot tugas '+sid,issues,True)
            linked_ihp=sorted({xid for x in risks if x.get('id') in risk_links for xid in x.get('ihp_ids',[])})
            objective_text='; '.join(oid+(' [usulan tentatif]' if objective_map[oid].get('source',{}).get('kind')=='USULAN_TENTATIF' else '')+': '+objective_map[oid].get('text','') for oid in links if oid in objective_map)
            tasks.append({'id':sid,'objective':objective_text+'\nPengujian: '+task.get('objective',''),'objective_ids':links,'risk_ids':risk_links,'ihp_ids':linked_ihp,'text':step.get('text',''),'assignees':names,'weight':weight,'reference':step.get('reference',''),'hours_by_person':{}})
    if not tasks or len(set(step_ids))!=len(step_ids):issues.append('KM 6: langkah tugas harus ada dan ID-nya unik.')
    if set(objective_map)-covered_objectives:issues.append('TAO/TEO belum dilaksanakan dalam KM 6: '+', '.join(sorted(set(objective_map)-covered_objectives)))
    split=plan.get('general_split',{});conclusion=number(split.get('conclusion_share',0),'Porsi jam simpulan',issues);final=number(split.get('final_kk_share',.5),'Porsi final kertas kerja',issues)
    if conclusion>=1 or final>1:issues.append('Porsi simpulan harus <1; porsi final KK harus <=1.')
    for name,p in people.items():
        person_tasks=[s for s in tasks if name in s['assignees']]
        if p['hours'][1]>0 and not person_tasks:issues.append('KM 6: personel beranggaran pelaksanaan belum diberi task: '+name)
        if person_tasks and p['hours'][1]<=0:issues.append('KM 6: pelaksana task tidak memiliki anggaran pelaksanaan: '+name)
        budget=p['hours'][1]*(1-conclusion);weight=sum(s['weight'] for s in person_tasks);assigned=0
        for i,s in enumerate(person_tasks):
            allocation=round(budget-assigned,6) if i==len(person_tasks)-1 else round(budget*s['weight']/weight,6)
            s['hours_by_person'][name]=allocation;assigned+=allocation
    pia=plan.get('pia',{})
    if m['jenis']=='PENGAWASAN' and 9 not in m['files']:
        for key in PIA_REQUIRED:
            if not nonempty(pia.get(key)):issues.append('Hasil PIA sebelumnya belum tersedia: '+key)
        for k in ['NIP_KT_PIA','NIP_KORWAS_PIA']:
            if nonempty(pia.get(k)):
                clean_pia_nip=re.sub(r'\s+','',str(pia[k]).strip())
                if len(clean_pia_nip)!=18 or not clean_pia_nip.isdigit():issues.append(k+': NIP harus 18 digit angka.')
    if issues:raise InputError(issues)
    rows=[]
    for s in tasks:
        hours=sum(s['hours_by_person'].values())
        rows.append([s['id'],s['objective']+'\nRisiko: '+', '.join(s['risk_ids'])+'\nInsilwas: '+', '.join(s['ihp_ids']),s['text'],'\n'.join(s['assignees']),hours,hours/m['jam_day'],s['reference']])
    m.update({'narratives':narratives,'pia':pia if m['jenis']=='PENGAWASAN' else {},'simpulan':m['stages'][1]*conclusion,'final_kk':m['stages'][2]*final,'plan':plan,'traceability':tasks})
    m['ihp']=ihp;m['tables'].update({'IHP':[[x['information'],x.get('source','')] for x in ihp],'PROSES_BISNIS':processes,'RISIKO_KLIEN':risk_rows,'RISIKO_PENUGASAN':assignment,'PROGRAM_KERJA':rows})
    return m
