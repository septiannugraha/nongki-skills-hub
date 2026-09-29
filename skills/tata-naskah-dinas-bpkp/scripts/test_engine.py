#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Smoke Test Engine Tata Naskah Dinas BPKP
=========================================
Menguji otomatis fungsionalitas pemuatan template, swap profil perwakilan,
dan injeksi konten pada aneka jenis naskah dinas resmi BPKP.
"""

import os
import sys
import tempfile
import docx

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

from generate_naskah import generate_document, load_registry, load_profiles


def run_tests():
    print("=" * 60)
    print("Memulai Automated Smoke Test Engine Naskah Dinas BPKP")
    print("=" * 60)

    # 1. Cek Registry & Profiles
    registry = load_registry()
    profiles = load_profiles()
    assert len(registry) == 46, f"Harus ada 46 template, ditemukan {len(registry)}"
    assert len(profiles) >= 20, f"Harus ada minimal 20 profil unit kerja, ditemukan {len(profiles)}"
    print(f"[PASS] Registry ({len(registry)} template) & Unit Profiles ({len(profiles)} unit) OK.")

    # 2. Uji Berbagai Kategori Dokumen
    test_cases = [
        {
            "jenis": "surat_tugas",
            "unit": "papua_tengah",
            "data": {
                "nomor_urut": "15",
                "keperluan": "Evaluasi Akuntabilitas Keuangan Desa",
                "tim": [{"no": 1, "nama": "Penguji Satu", "nip": "19800101", "peran": "Ketua Tim"}]
            }
        },
        {
            "jenis": "nota_dinas",
            "unit": "kalsel",
            "data": {
                "nomor_urut": "77",
                "replacements": {"Hal\t:": "Hal\t: Undangan Rapat Kerja Pimpinan"}
            }
        },
        {
            "jenis": "surat_dinas",
            "unit": "jatim",
            "data": {
                "nomor_urut": "99",
                "replacements": {"Hal\t:": "Hal\t: Permintaan Bahan Evaluasi"}
            }
        },
        {
            "jenis": "berita_acara",
            "unit": "papua",
            "data": {
                "nomor_urut": "03",
                "keperluan": "Serah Terima Hasil Pengawasan"
            }
        },
        {
            "jenis": "keputusan_kepala",
            "unit": "pusat_utama",
            "data": {
                "nomor_urut": "500",
                "keperluan": "Penetapan Standar Pengawasan"
            }
        }
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        for idx, tc in enumerate(test_cases, 1):
            out_docx = os.path.join(tmpdir, f"test_{tc['jenis']}_{tc['unit']}.docx")
            res_path = generate_document(
                template_key=tc["jenis"],
                unit_key=tc["unit"],
                data=tc["data"],
                output_path=out_docx,
                kop_mode="table" if tc["jenis"] != "keputusan_kepala" else "text"
            )
            assert os.path.exists(res_path), f"File {res_path} gagal dibuat"
            
            # Verifikasi keterbacaan python-docx
            doc = docx.Document(res_path)
            assert len(doc.paragraphs) > 0 or len(doc.tables) > 0, "Dokumen kosong!"
            print(f"[PASS] Case {idx}: {tc['jenis']} ({tc['unit']}) -> Generated successfully ({os.path.getsize(res_path)} bytes).")

    print("=" * 60)
    print("SEMUA AUTOMATED TESTS BERHASIL DILALUI DENGAN SUKSES! (100% OK)")
    print("=" * 60)


if __name__ == "__main__":
    run_tests()
