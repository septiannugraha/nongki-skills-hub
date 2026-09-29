#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Unified Naskah Dinas Generator - BPKP
====================================
Engine utama untuk menyusun naskah dinas resmi BPKP berbasis master template Word (.docx)
sesuai Peraturan BPKP Nomor 4 Tahun 2026.

Fitur Unggulan:
1. Mendukung 46 Master Template Dokumen Dinas BPKP (assets/templates/*.docx).
2. Sistem Unit Profiles (assets/unit_profiles.json): Kop surat, kode unit, dan kota
   otomatis berganti antarperwakilan (misal Papua Tengah vs Kalsel vs Pusat).
3. Dua mode Kop Surat: In-Place Text Replacement atau Official Table Kop (2 kolom + logo BPKP).
4. Penomoran otomatis sesuai taksonomi Peraturan BPKP 4/2026.
5. Injeksi baris tabel dinamis (misal daftar tim personil surat tugas).

Usage CLI:
    python generate_naskah.py --list
    python generate_naskah.py --jenis surat_tugas --unit papua_tengah --output surat_tugas_pt.docx
    python generate_naskah.py --jenis surat_tugas --unit kalsel --data data.json --output surat_tugas_kalsel.docx
"""

from __future__ import annotations

import os
import sys
import json
import re
import argparse
from datetime import datetime
from typing import Dict, Any, List, Optional

import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Path konfigurasi default
_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_TEMPLATES_DIR = os.path.join(_BASE_DIR, "assets", "templates")
_PROFILES_FILE = os.path.join(_BASE_DIR, "assets", "unit_profiles.json")
_REGISTRY_FILE = os.path.join(_BASE_DIR, "references", "templates_registry.json")
_LOGO_PATH = os.path.join(_BASE_DIR, "assets", "logo_bpkp.png")


def load_profiles() -> Dict[str, Any]:
    """Muat seluruh data profil unit kerja BPKP."""
    if os.path.exists(_PROFILES_FILE):
        with open(_PROFILES_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def load_registry() -> Dict[str, Any]:
    """Muat registry katalog ke-46 template."""
    if os.path.exists(_REGISTRY_FILE):
        with open(_REGISTRY_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def get_profile(profile_key: str, overrides: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Dapatkan profil unit kerja dengan opsi fallback dan override."""
    profiles = load_profiles()
    key = profile_key.lower().replace("-", "_").replace(" ", "_")
    profile = profiles.get(key)
    
    if not profile:
        # Fallback pencarian parsial
        for k, v in profiles.items():
            if key in k or k in key:
                profile = v.copy()
                break
                
    if not profile:
        # Profil default jika tidak ditemukan
        profile = {
            "kode_unit": "PW",
            "nama_kantor": f"PERWAKILAN BPKP {profile_key.upper()}",
            "nama_kop_baris1": "BADAN PENGAWASAN KEUANGAN DAN PEMBANGUNAN",
            "nama_kop_baris2": f"PERWAKILAN PROVINSI {profile_key.upper()}",
            "alamat": "Alamat Kantor Perwakilan BPKP",
            "wilayah_kodepos": "Kode Pos -",
            "telepon": "-",
            "faksimile": "-",
            "email": f"{key}@bpkp.go.id",
            "website": f"www.bpkp.go.id/{key}",
            "kota": profile_key.title(),
            "jabatan_kepala": "Kepala Perwakilan"
        }
    else:
        profile = profile.copy()

    if overrides:
        profile.update(overrides)
    return profile


def replace_text_in_paragraph(paragraph, old_text: str, new_text: str) -> bool:
    """Ganti teks dalam satu paragraf dengan tetap mempertahankan run format jika memungkinkan."""
    if old_text not in paragraph.text:
        return False
        
    # Jika old_text persis berada dalam 1 run:
    replaced = False
    for run in paragraph.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            replaced = True
            
    if replaced:
        return True
        
    # Jika old_text terpecah di antara beberapa runs (kompleks), gabungkan sementara
    full_text = paragraph.text
    new_full = full_text.replace(old_text, new_text)
    if paragraph.runs:
        paragraph.runs[0].text = new_full
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.text = new_full
    return True


def replace_text_in_table(table, old_text: str, new_text: str) -> int:
    """Ganti teks di seluruh sel dalam tabel."""
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if replace_text_in_paragraph(p, old_text, new_text):
                    count += 1
    return count


def replace_all_placeholders(doc: docx.Document, mapping: Dict[str, str]):
    """Gantikan semua dictionary placeholder di seluruh paragraf dan tabel docx."""
    for old_k, new_v in mapping.items():
        if new_v is None:
            continue
        str_val = str(new_v)
        for p in doc.paragraphs:
            replace_text_in_paragraph(p, old_k, str_val)
        for t in doc.tables:
            replace_text_in_table(t, old_k, str_val)


def apply_unit_kop_text(doc: docx.Document, profile: Dict[str, Any]):
    """
    Terapkan profil unit kerja pada format kop teks bawaan template:
      ... (UNIT KERJA) -> PERWAKILAN PROVINSI ...
      ... (alamat) -> Jalan ...
      Telepon ... -> Telepon ...
      E-mail ... -> Email ...
    """
    for p in doc.paragraphs[:8]:
        t = p.text
        if "... (UNIT KERJA)" in t or "… (UNIT KERJA)" in t:
            p.text = profile.get("nama_kop_baris2", "")
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(12)
                run.font.bold = True
        elif "... (alamat)" in t or "… (alamat)" in t:
            p.text = profile.get("alamat", "")
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
        elif "nomor telepon" in t:
            telp = profile.get("telepon", "-")
            faks = profile.get("faksimile", "-")
            p.text = f"Telepon {telp}, Faksimile {faks}"
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
        elif "Website www.bpkp.go.id" in t:
            em = profile.get("email", "")
            web = profile.get("website", "www.bpkp.go.id")
            p.text = f"E-mail {em}, Website {web}"
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)


def insert_official_kop_table(doc: docx.Document, profile: Dict[str, Any], logo_path: str = _LOGO_PATH):
    """
    Gantikan kop teks bawaan dengan Tabel Kop Resmi 2 kolom + Logo BPKP + Garis Tebal Bawah
    sesuai standar Peraturan BPKP 4/2026.
    """
    # 1. Deteksi dan bersihkan 5 baris kop teks di awal dokumen
    del_indices = []
    for i, p in enumerate(doc.paragraphs[:7]):
        txt = p.text.strip()
        if any(keyword in txt for keyword in [
            "BADAN PENGAWASAN KEUANGAN DAN PEMBANGUNAN",
            "UNIT KERJA",
            "nomor telepon",
            "Website www.bpkp.go.id"
        ]):
            del_indices.append(p)
            
    for p in del_indices:
        p.text = ""

    # 2. Sisipkan tabel kop di posisi paling atas
    # Buat tabel 1 row x 2 cols
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Borderless kecuali border bawah tebal (sz=18 / 2.25pt)
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="18" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    # Sel Kiri: Logo BPKP
    cell_logo = table.cell(0, 0)
    cell_logo.width = Cm(3.2)
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(4)
    if os.path.exists(logo_path):
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Cm(2.6))

    # Sel Kanan: Teks Instansi Resmi
    cell_text = table.cell(0, 1)
    cell_text.width = Cm(13.3)
    p_text = cell_text.paragraphs[0]
    p_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_text.paragraph_format.space_before = Pt(0)
    p_text.paragraph_format.space_after = Pt(0)
    p_text.paragraph_format.line_spacing = 1.0

    # Baris 1: INDUK
    r1 = p_text.add_run("BADAN PENGAWASAN KEUANGAN DAN PEMBANGUNAN\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.bold = True

    # Baris 2: UNIT / PERWAKILAN
    baris2 = profile.get("nama_kop_baris2", profile.get("nama_kantor", ""))
    if baris2:
        r2 = p_text.add_run(f"{baris2.upper()}\n")
        r2.font.name = "Arial"
        r2.font.size = Pt(12)
        r2.font.bold = True

    # Baris 3: Alamat
    alamat = profile.get("alamat", "")
    if alamat:
        r3 = p_text.add_run(f"{alamat}\n")
        r3.font.name = "Arial"
        r3.font.size = Pt(9)

    # Baris 4: Wilayah & Kode Pos
    wilayah = profile.get("wilayah_kodepos", "")
    if wilayah:
        r4 = p_text.add_run(f"{wilayah}\n")
        r4.font.name = "Arial"
        r4.font.size = Pt(9)

    # Baris 5: Kontak & Website
    telp = profile.get("telepon", "-")
    em = profile.get("email", "")
    web = profile.get("website", "www.bpkp.go.id")
    r5 = p_text.add_run(f"Telepon: {telp} | E-mail: {em} | Website: {web}")
    r5.font.name = "Arial"
    r5.font.size = Pt(8.5)

    # Pindahkan tabel ke awal dokumen di XML body
    doc._body._element.insert(0, table._tbl)


def populate_personel_table(doc: docx.Document, personel_list: List[Dict[str, str]]):
    """
    Isi tabel personil tim pada Surat Tugas.
    Mencari tabel yang memiliki header No / Nama / NIP / Jabatan/Peran.
    """
    if not personel_list:
        return
        
    for table in doc.tables:
        header_text = " ".join([c.text.lower() for c in table.rows[0].cells])
        if "nama" in header_text and ("nip" in header_text or "peran" in header_text or "jabatan" in header_text):
            # Baris template personil ada di index 1
            # Hapus baris placeholder default jika ada
            while len(table.rows) > 1:
                table._tbl.remove(table.rows[1]._tr)
                
            # Tambahkan baris personil aktual
            for i, p in enumerate(personel_list, 1):
                row = table.add_row()
                cells = row.cells
                cells[0].text = str(p.get("no", i))
                cells[1].text = p.get("nama", "-")
                cells[2].text = str(p.get("nip", "-"))
                cells[3].text = p.get("peran", p.get("jabatan", "-"))
                
                # Format font Arial 11 pt
                for cell in cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(2)
                        for r in cp.runs:
                            r.font.name = "Arial"
                            r.font.size = Pt(11)
            break


def generate_document(template_key: str,
                      unit_key: str = "papua_tengah",
                      data: Optional[Dict[str, Any]] = None,
                      output_path: Optional[str] = None,
                      kop_mode: str = "table") -> str:
    """
    Generator utama: memuat master template docx, menerapkan profil kop unit,
    menginjeksi data payload, dan menghasilkan naskah dinas Word final.
    """
    registry = load_registry()
    key = template_key.lower().replace("-", "_").replace(" ", "_")
    
    if key not in registry:
        # Cari template yang mendekati
        matched = [k for k in registry if key in k]
        if matched:
            key = matched[0]
        else:
            raise ValueError(f"Template '{template_key}' tidak ditemukan dalam registry. Gunakan --list untuk melihat.")

    tpl_info = registry[key]
    tpl_file = os.path.join(_TEMPLATES_DIR, tpl_info["filename"])
    if not os.path.exists(tpl_file):
        raise FileNotFoundError(f"File template tidak ditemukan: {tpl_file}")

    # Muat dokumen
    doc = docx.Document(tpl_file)
    profile = get_profile(unit_key)
    payload = data or {}

    # 1. Tangani Kop Surat
    if tpl_info.get("has_unit_kop", False):
        if kop_mode == "table":
            insert_official_kop_table(doc, profile, _LOGO_PATH)
        else:
            apply_unit_kop_text(doc, profile)

    # 2. Siapkan Kamus Penggantian Cerdas (Placeholders)
    now = datetime.now()
    thn = str(payload.get("tahun", now.year))
    bln_indo = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
                "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    tgl_indo = f"{now.day} {bln_indo[now.month - 1]} {thn}"

    # Format nomor naskah
    kd_arsip = payload.get("kode_klasifikasi", "PW.01.03")
    no_urut = payload.get("nomor_urut", "101")
    kd_unit = payload.get("kode_unit", profile.get("kode_unit", "PW34"))
    konseptor = payload.get("konseptor", "2")
    kd_naskah = tpl_info.get("kode_naskah", "ND")

    formatted_nomor = f"{kd_arsip}/{kd_naskah}-{no_urut}/{kd_unit}/{konseptor}/{thn}"
    if payload.get("nomor_lengkap"):
        formatted_nomor = payload["nomor_lengkap"]

    kota = profile.get("kota", "Nabire")
    tanggal_naskah = payload.get("tanggal_naskah", f"{kota}, {tgl_indo}")

    replacements = {
        # Pola umum kurung di template
        "... (UNIT KERJA)": profile.get("nama_kop_baris2", ""),
        "… (UNIT KERJA)": profile.get("nama_kop_baris2", ""),
        "... (alamat)": profile.get("alamat", ""),
        "… (alamat)": profile.get("alamat", ""),
        "... (tanggal, bulan, dan tahun)": tanggal_naskah,
        "… (tanggal, bulan, dan tahun)": tanggal_naskah,
        "... (tanggal, bulan dan tahun)": tanggal_naskah,
        "... (nama tempat)": kota,
        "… (nama tempat)": kota,
        "... (tempat)": kota,
        "... (kota)": kota,
        "... (tahun pembuatan)": thn,
        "… (tahun pembuatan)": thn,
        "... (kode unit atau jabatan)": kd_unit,
        "... (konseptor naskah dinas)": str(konseptor),
        "... (nomor urut naskah dinas)": str(no_urut),
        "... (kode klasifikasi arsip)": kd_arsip,
        "... (pimpinan unit kerja)": payload.get("pimpinan_unit_kerja", profile.get("jabatan_kepala", "Kepala Perwakilan")),
        "(nama penugasan)": payload.get("nama_penugasan", payload.get("keperluan", "")),
        "... (unit kerja)": profile.get("nama_kantor", ""),
        "… (unit kerja)": profile.get("nama_kantor", ""),
        "... (nama jabatan)": payload.get("jabatan_penandatangan", profile.get("jabatan_kepala", "Kepala Perwakilan")),
        "... (nama lengkap)": payload.get("nama_penandatangan", "Pejabat Penanda Tangan"),
        "(narasi muatan/isi Surat Dinas)": payload.get("isi_surat", payload.get("isi", "")),
        "(narasi muatan/isi Nota Dinas)": payload.get("isi_nota_dinas", payload.get("isi", ""))
    }

    # Tambahkan custom mapping dari payload
    if "replacements" in payload:
        replacements.update(payload["replacements"])

    # Jalankan penggantian teks placeholder
    replace_all_placeholders(doc, replacements)

    # Khusus nomor surat/naskah jika ada pola regex NOMOR ...
    for p in doc.paragraphs[:15]:
        if "NOMOR" in p.text and ("kode klasifikasi" in p.text or "nomor urut" in p.text):
            p.text = f"NOMOR {formatted_nomor}"
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11)
                r.font.bold = True

    # 3. Tangani Injeksi Tabel Personil (untuk Surat Tugas)
    if "personel" in payload or "tim" in payload:
        tim = payload.get("personel") or payload.get("tim")
        populate_personel_table(doc, tim)

    # 4. Tentukan output path
    if not output_path:
        out_name = f"{key}_{unit_key}_{thn}.docx"
        output_path = os.path.join(os.getcwd(), out_name)

    doc.save(output_path)
    return output_path


def print_list_templates():
    """Tampilkan daftar seluruh 46 template."""
    reg = load_registry()
    print(f"\n{'='*75}")
    print(f"{'DAFTAR 46 MASTER TEMPLATE NASKAH DINAS BPKP (PERATURAN 4/2026)':^75}")
    print(f"{'='*75}")
    print(f"{'ID':<30} | {'KODE':<6} | {'KATEGORI':<25}")
    print(f"{'-'*75}")
    for k, v in sorted(reg.items()):
        print(f"{k:<30} | {v.get('kode_naskah','-'):<6} | {v.get('kategori','-')[:25]:<25}")
    print(f"{'='*75}")
    print(f"Total: {len(reg)} template tersedia di assets/templates/\n")


def print_list_profiles():
    """Tampilkan daftar unit profiles yang tersimpan."""
    profiles = load_profiles()
    print(f"\n{'='*75}")
    print(f"{'DAFTAR PROFIL UNIT KERJA / KOP PERWAKILAN BPKP':^75}")
    print(f"{'='*75}")
    print(f"{'KEY':<18} | {'KODE':<6} | {'NAMA UNIT / KANTOR':<35} | {'KOTA':<10}")
    print(f"{'-'*75}")
    for k, v in sorted(profiles.items()):
        print(f"{k:<18} | {v.get('kode_unit','-'):<6} | {v.get('nama_kantor','-')[:35]:<35} | {v.get('kota','-'):<10}")
    print(f"{'='*75}\n")


def main():
    parser = argparse.ArgumentParser(description="Unified Naskah Dinas BPKP Generator")
    parser.add_argument("--list", action="store_true", help="Tampilkan semua 46 template yang tersedia")
    parser.add_argument("--list-profiles", action="store_true", help="Tampilkan daftar profil kop unit kerja yang tersimpan")
    parser.add_argument("--jenis", type=str, help="ID template naskah dinas (misal: surat_tugas, nota_dinas, berita_acara)")
    parser.add_argument("--unit", type=str, default="papua_tengah", help="Kunci profil unit kerja (default: papua_tengah, kalsel, jatim, dki, dll.)")
    parser.add_argument("--data", type=str, help="Path ke file JSON data masukan")
    parser.add_argument("--output", type=str, help="Path file Word .docx hasil generate")
    parser.add_argument("--kop-mode", type=str, choices=["table", "text"], default="table",
                        help="Mode kop surat: 'table' (resmi 2-kolom logo BPKP) atau 'text' (teks in-place)")

    args = parser.parse_args()

    if args.list:
        print_list_templates()
        return

    if args.list_profiles:
        print_list_profiles()
        return

    if not args.jenis:
        parser.print_help()
        return

    payload = {}
    if args.data and os.path.exists(args.data):
        with open(args.data, "r", encoding="utf-8") as f:
            payload = json.load(f)

    out_file = generate_document(
        template_key=args.jenis,
        unit_key=args.unit,
        data=payload,
        output_path=args.output,
        kop_mode=args.kop_mode
    )
    print(f"[OK] Sukses menyusun naskah dinas: {out_file}")


if __name__ == "__main__":
    main()
